import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def md_to_docx(md_path, docx_path):
    doc = docx.Document()

    # Define standard margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B) # Slate-800

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    def flush_table(lines):
        if not lines:
            return
        rows = [l.strip().strip('|').split('|') for l in lines if '|' in l and not '---' in l]
        if not rows:
            return
        
        cols_count = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                if c_idx < cols_count:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = val.strip()
                    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                    
                    if r_idx == 0:
                        set_cell_background(cell, "0F172A") # Slate-900
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                run.font.size = Pt(9.5)
                    else:
                        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg_color)
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(9.5)
                                run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        doc.add_paragraph() # Spacing

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')
        
        # Table Detection
        if '|' in line and not line.startswith('>'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                flush_table(table_lines)
                in_table = False
                table_lines = []

        if not line.strip():
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[2:].strip())
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Deep Slate
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:].strip())
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF) # Blue-800
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[4:].strip())
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x3B, 0x82, 0xF6) # Blue-500
        elif line.startswith('> '):
            # Callout block
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[2:].strip())
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            # Parse bold runs
            text = line[2:].strip()
            parts = text.split('**')
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                if idx % 2 == 1:
                    run.font.bold = True
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            text = line.strip()
            parts = text.split('**')
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                if idx % 2 == 1:
                    run.font.bold = True

        i += 1

    if in_table:
        flush_table(table_lines)

    doc.save(docx_path)
    print(f"Successfully generated DOCX at {docx_path}")

if __name__ == "__main__":
    md_file = r"e:\COLLEGE\Hackathon\SAS-CU-HACKATHON\SAS_CU_Hackathon_The_Processors_Final_Approach_Note.md"
    docx_file = r"e:\COLLEGE\Hackathon\SAS-CU-HACKATHON\SAS_CU_Hackathon_The_Processors_Final_Approach_Note.docx"
    md_to_docx(md_file, docx_file)
